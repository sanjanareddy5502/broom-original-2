"""Regression tests focus on preservation, approval, imports, and adversarial data."""
from io import BytesIO
import json
import zipfile
import random

import pandas as pd
import pytest
from openpyxl import load_workbook

from broom_core.security import read_uploaded_file, InputRejected, MAX_UPLOAD_BYTES
from broom_core.pipeline import diagnose, execute, validate_plan, plan_signature
from broom_core.recipes import export_recipe, import_recipe
from numeric_intelligence import parse_number, parse_series
from date_intelligence import inspect_dates
from business_brief import build_business_brief
from cleaner import clean_dataset
from builder import _spreadsheet_safe, build_outputs


def test_csv_preserves_leading_zeros_and_codes():
    result = read_uploaded_file(b'postal_code,invoice,amount\n00101,00023,"$1,200"\n', 'example.csv')
    assert result.iloc[0].tolist() == ['00101', '00023', '$1,200']


def test_csv_rejects_bad_records():
    with pytest.raises(InputRejected):
        read_uploaded_file(b'a,b\n1,2\n"unclosed,3\n', 'bad.csv')


def test_csv_rejects_null_binary_and_wrong_extension():
    for name, data in [('a.csv',b'a,b\x00,1'), ('a.exe', b'a,b\n1,2')]:
        with pytest.raises(InputRejected):
            read_uploaded_file(data, name)


def test_upload_size_and_empty_limits():
    for data in [b'', b'a' * (MAX_UPLOAD_BYTES + 1)]:
        with pytest.raises(InputRejected):
            read_uploaded_file(data, 'a.csv')


def test_xlsx_rejects_fake_zip():
    with pytest.raises(InputRejected):
        read_uploaded_file(b'not an excel workbook', 'example.xlsx')


def test_xlsx_rejects_decompression_attack():
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', 'x')
        archive.writestr('xl/workbook.xml', 'x')
        archive.writestr('xl/worksheets/sheet1.xml', '0' * (26 * 1024 * 1024))
    with pytest.raises(InputRejected):
        read_uploaded_file(buffer.getvalue(), 'compressed.xlsx')


def test_xlsx_good_file_imports():
    buffer = BytesIO()
    pd.DataFrame({'name':['ok'], 'amount':[30]}).to_excel(buffer, index=False)
    result = read_uploaded_file(buffer.getvalue(), 'safe.xlsx')
    assert result.loc[0, 'amount'] == 30


@pytest.mark.parametrize('raw', ['--', 'NA', '12,34', '12%', '00091', '1,2,300', '1e309', '$12USD', 'not known'])
def test_numeric_invalid_formats_fail_closed(raw):
    assert pd.isna(parse_number(raw))


@pytest.mark.parametrize('raw, expected', [('$1,200.00', 1200), ('($1,200.00)', -1200), ('-$1.20', -1.2), ('0.20', 0.2), ('1,234,567', 1234567), ('42', 42)])
def test_numeric_valid_formats(raw, expected):
    assert parse_number(raw) == expected


def test_identifiers_not_turned_into_numbers_even_when_requested():
    df = pd.DataFrame({'postal_code':['00101','00102']})
    out, log = clean_dataset(df, [{'action':'convert_text_to_numeric','approved':True,'columns':['postal_code']}])
    assert out.equals(df)


def test_conflicting_date_evidence_is_review_not_inferred():
    result = inspect_dates(pd.Series(['15/04/2026','04/15/2026','01/02/2026']))
    assert result['status'] == 'review'
    assert result['ambiguous_count'] == 3


def test_user_confirmed_date_preference_still_rejects_wrong_format():
    result = inspect_dates(pd.Series(['15/04/2026','04/15/2026']), 'day_first')
    assert result['status'] == 'review'
    assert result['invalid_count'] == 1


def test_mixed_iso_and_day_first_dates_with_timings():
    result = inspect_dates(pd.Series(['2026-02-01 15:30:00','15/02/2026', '04/05/2026']))
    assert result['status'] == 'ready'
    assert result['parsed'].iloc[0].hour == 15
    assert result['parsed'].iloc[2].month == 5


def test_duplicate_dataframe_index_does_not_overwrite_wrong_row():
    s = pd.Series(['2026-02-01', '2026-03-01'], index=[7,7])
    result = inspect_dates(s)
    assert result['status'] == 'ready'
    assert result['parsed'].iloc[0].month == 2
    assert result['parsed'].iloc[1].month == 3


def test_date_with_no_year_never_invented():
    assert inspect_dates(pd.Series(['March 3', 'May 5']))['status'] == 'review'


def test_invalid_date_never_overwrites_good_record():
    data = pd.DataFrame({'event_date':['2026-02-01','not a date']})
    out, audit = clean_dataset(data, [{'action':'convert_text_to_datetime','approved':True,'columns':['event_date']}])
    assert out.equals(data)
    assert any(x['action'] == 'needs_date_review' for x in audit['cleaning_log'])


def test_categories_preserve_capitalization_and_reconcile_variants():
    df = pd.DataFrame({'location':['USA','USA','usa','UK','UK']})
    after, _ = clean_dataset(df,[{'action':'standardize_categories','approved':True,'columns':['location']}])
    assert after['location'].tolist() == ['USA','USA','USA','UK','UK']


def test_category_edit_not_approved_by_default():
    d = diagnose(pd.DataFrame({'category':['USA','usa','USA']}))
    plan = [p for p in d.plan if p['action'] == 'standardize_categories']
    assert len(plan) == 1 and plan[0]['requires_approval'] and not plan[0]['approved']


def test_preserve_hyphens_and_NA_business_values():
    df = pd.DataFrame({'region':['NA','-','  CA  ',None]})
    cleaned, _ = clean_dataset(df, [{'action':'trim_text_and_standardize_nulls','approved':True}])
    assert cleaned['region'].tolist()[:3] == ['NA','-', 'CA']


def test_flag_columns_cannot_silently_overwrite_existing():
    df = pd.DataFrame({'Amount':[1,2,3,4,100], 'datatrust_amount_outlier_flag':['preserve']*5})
    after, log = clean_dataset(df,[{'action':'flag_outliers','approved':True,'columns':['Amount']}])
    assert after['datatrust_amount_outlier_flag'].tolist() == ['preserve']*5
    assert any(e['action'] == 'skipped_flag_name_collision' for e in log['cleaning_log'])


def test_duplicate_transaction_flag_collision_is_fatal():
    df = pd.DataFrame({'name':['A','A'], 'datatrust_potential_duplicate_transaction':[None,None]})
    with pytest.raises(ValueError, match='flag column'):
        clean_dataset(df,[{'action':'flag_duplicate_transactions','approved':True}])


def test_pipeline_rejects_unapproved_high_risk_action():
    with pytest.raises(ValueError,match='Explicit confirmation'):
        validate_plan([{'action':'remove_exact_duplicate_rows','approved':True}])
    validate_plan([{'action':'remove_exact_duplicate_rows','approved':True,'approval_confirmed':True}])


def test_pipeline_rejects_unknown_action_and_date_preference():
    with pytest.raises(ValueError):
        validate_plan([{'action':'run_untrusted_python','approved':True}])
    with pytest.raises(ValueError):
        validate_plan([{'action':'convert_text_to_datetime','approved':True,'approval_confirmed':True,
                        'metadata':{'date_preference':'guess_anything'}}])


def test_plan_fingerprint_changes_on_preference_or_approval():
    a = [{'action':'convert_text_to_datetime','approved':False,'columns':['date']}]
    b = [{**a[0], 'approved':True, 'approval_confirmed':True}]
    assert plan_signature(a) != plan_signature(b)


def test_imported_recipe_does_not_auto_approve_destructive_changes():
    original = [{'action':'remove_exact_duplicate_rows','columns':[],'approved':False},
                {'action':'convert_text_to_datetime','columns':['date'],'approved':False,'metadata':{}}]
    incoming = [{**original[0], 'approved':True},
                {**original[1], 'approved':True,'metadata':{'date_preference':'day_first'}}]
    restored = import_recipe(export_recipe(incoming).encode(), original)
    assert not restored[0]['approved'] and not restored[1]['approved']
    assert restored[1]['metadata']['date_preference'] == 'day_first'


def test_imported_recipe_rejects_oversized_or_invalid():
    plan = [{'action':'convert_text_to_datetime','columns':['date'],'approved':False}]
    for content in [b'not json', b'0'*(65536+1), b'{"version": 4, "preferences": []}']:
        with pytest.raises(ValueError):
            import_recipe(content, plan)


def test_pipeline_does_not_mutate_raw_upload():
    df = pd.DataFrame({'customer':['A','A'], 'price':['$2','3']})
    initial = df.copy(deep=True)
    d = diagnose(df)
    result = execute(df, d, d.plan)
    pd.testing.assert_frame_equal(df, initial)
    assert result.outputs['excel'].getbuffer().nbytes > 1000


def test_bi_monthly_uses_actual_dated_records():
    df = pd.DataFrame({'order_date':pd.to_datetime(['2026-01-01', '2026-02-01', None]),
                       'amount':[10.,20.,30.]})
    result = build_business_brief(df)
    assert result['metrics']['undated_records'] == 1
    assert result['monthly']['records'].tolist() == [1,1]
    assert result['monthly']['measure_total'].tolist() == [10,20]
    assert any('cannot be placed' in q for q in result['questions'])


def test_excel_csv_formula_injection_and_original_preserved():
    df = pd.DataFrame({'danger':['=1+1', '+1+1', '\t@SUM(1,2)', '-CMD()', 'normal']})
    safe = _spreadsheet_safe(df)
    assert df['danger'].iloc[0] == '=1+1'
    assert all(x.startswith("'") for x in safe['danger'].iloc[:4])
    d = diagnose(df)
    output = execute(df, d, d.plan).outputs
    wb = load_workbook(BytesIO(output['excel'].getvalue()))
    assert all(wb['cleaned_data'].cell(row=i,column=1).data_type != 'f' for i in range(2,7))


def test_randomized_date_round_trip_300_cases():
    rng = random.Random(13)
    raw = []
    for _ in range(300):
        month = rng.randint(1,12)
        day = rng.randint(13,28)
        raw.append(f'{day:02d}/{month:02d}/2026')
    result = inspect_dates(pd.Series(raw))
    assert result['status'] == 'ready'
    for before, after in zip(raw, result['parsed']):
        day, month, year = (int(x) for x in before.split('/'))
        assert (after.day,after.month,after.year) == (day,month,year)


def test_relationship_intelligence_flags_mismatch_without_editing():
    df = pd.DataFrame({'quantity':[2,3], 'unit_price':[10,10], 'line_total':[20,25]})
    brief = build_business_brief(df)
    relationship = brief['relationships'][0]
    assert relationship['checked_rows'] == 2
    assert relationship['mismatched_rows'] == 1
    assert df.loc[1,'line_total'] == 25


def test_relationship_intelligence_allows_rounding():
    df = pd.DataFrame({'quantity':[3], 'unit_price':[0.3333333], 'line_total':[1.00]})
    relationship = build_business_brief(df)['relationships'][0]
    assert relationship['mismatched_rows'] == 0


def test_xlsx_formula_cannot_be_silently_read_as_missing():
    from openpyxl import Workbook
    workbook = Workbook()
    ws = workbook.active
    ws['A1'] = 'revenue'
    ws['A2'] = '=10+10'
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(InputRejected, match='formulas'):
        read_uploaded_file(buffer.getvalue(), 'formula.xlsx')


def test_nonfinite_numbers_are_not_accepted():
    assert pd.isna(parse_number(float('inf')))
    assert pd.isna(parse_number(float('-inf')))
