"""Repeatable synthetic end-to-end regression scenarios across six business-data shapes."""
from io import BytesIO
import random
import pandas as pd
import pytest
from openpyxl import load_workbook

from broom_core.pipeline import diagnose, execute


@pytest.mark.parametrize('seed', range(120))
def test_generated_business_workflows_are_bounded_and_nondestructive(seed):
    rng = random.Random(seed)
    n = 8 + rng.randrange(8)
    if seed % 6 == 0:
        df = pd.DataFrame({'order_date':[f'2026-03-{(i%25)+1:02d}' for i in range(n)],
                           'amount':[f'${rng.randrange(1,800):,}' for _ in range(n)],
                           'order_id':[f'000{i+1:03d}' for i in range(n)]})
    elif seed % 6 == 1:
        df = pd.DataFrame({'vendor':['A','B','A','B'] * (n//4) + ['A']*(n%4),
                           'invoice':['INV-'+str(i) for i in range(n)],
                           'debit':[str(rng.randrange(1,500)) for _ in range(n)]})
    elif seed % 6 == 2:
        df = pd.DataFrame({'date':['01/02/2026' if i%2 else '02/03/2026' for i in range(n)],
                           'rate':['12%' if i%3 else '10%' for i in range(n)],
                           'country':['NA' if i%3 else 'US' for i in range(n)]})
    elif seed % 6 == 3:
        df = pd.DataFrame({'customer':['A','B','C']*(n//3)+['A']*(n%3),
                           'comments':['=SUM(1,2)' if i%3 else '-cmd()' for i in range(n)],
                           'phone':[f'00{i:04d}' for i in range(n)]})
    elif seed % 6 == 4:
        df = pd.DataFrame({'quantity':[rng.randrange(1,10) for _ in range(n)],
                           'unit_price':[rng.randrange(1,10) for _ in range(n)],
                           'line_total':[rng.randrange(1,100) for _ in range(n)]})
    else:
        df = pd.DataFrame({'location':['USA','usa','UK','uk']*(n//4)+['USA']*(n%4),
                           'misc':[None]*n,
                           'notes':['  task  ' if i%2 else 'ok' for i in range(n)]})
    raw = df.copy(deep=True)
    diagnosis = diagnose(df)
    result = execute(df, diagnosis, diagnosis.plan)
    pd.testing.assert_frame_equal(df, raw)
    assert result.validation['status']
    assert isinstance(result.summary['cleaning_log'], list)
    assert result.cleaned.shape[0] <= df.shape[0]
    output = result.outputs['excel'].getvalue()
    assert len(output) > 1000
    workbook = load_workbook(BytesIO(output), read_only=True)
    assert 'cleaned_data' in workbook.sheetnames
    workbook.close()
